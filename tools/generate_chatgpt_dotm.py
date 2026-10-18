import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import zipfile
import io
import os

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Segoe UI'
    normal_font.size = Pt(10.5)
    normal_font.color.rgb = RGBColor(0x1B, 0x1F, 0x2A) # Ink token

    # Document Header / Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = title.add_run("SRI VISHNU TEJA — COMPLETE PROJECT BRIEF FOR CHATGPT")
    run_title.font.name = 'Segoe UI Semibold'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x26, 0x40) # Navy token

    subtitle = doc.add_paragraph()
    run_sub = subtitle.add_run("End-to-End Architecture, Product Catalogue, Technical Specifications, and Development Roadmap")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x6B, 0x72, 0x80) # Muted token

    doc.add_paragraph("─" * 65)

    # ChatGPT Role Prompt Callout Box
    p_box = doc.add_paragraph()
    run_box_label = p_box.add_run("HOW TO USE THIS WITH CHATGPT:\n")
    run_box_label.bold = True
    run_box_label.font.color.rgb = RGBColor(0x8A, 0x52, 0x12) # Saffron text
    run_box_text = p_box.add_run(
        "Copy and paste this document or upload this file into ChatGPT. Use the prompt below:\n"
        "\"You are an expert senior full-stack software engineer and system architect. Acting as my pair programmer on this existing project, read this comprehensive end-to-end specification for 'Sri Vishnu Teja Books, Xerox & Cool Drinks'. Follow all architectural constraints (₹0 budget, counter-pickup only, strict design tokens, React 18/TypeScript/Firebase/Supabase stack). Acknowledge receipt and let me know you are ready to assist with the next development sprint.\""
    )
    run_box_text.font.italic = True

    doc.add_paragraph()

    # Section 1: Executive Overview & Business Model
    h1 = doc.add_heading("1. Executive Summary & Business Operations", level=1)
    h1.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    doc.add_paragraph(
        "Sri Vishnu Teja Books, Xerox & Cool Drinks is an established brick-and-mortar retail business in Railway Kodur, Andhra Pradesh, India. "
        "The web application is designed to solve in-store congestion and long counter queues by allowing school and college students, office workers, "
        "and local residents to browse inventory, order stationery, and submit Xerox printing jobs remotely from their mobile or desktop browsers."
    )

    p_facts = doc.add_paragraph()
    p_facts.add_run("Key Business & Operational Facts:\n").bold = True
    facts = [
        ("Business Name", "Sri Vishnu Teja Books, Xerox & Cool Drinks"),
        ("Proprietor", "Gevini Murali Mohan"),
        ("Location", "Near Bajaj Showroom, RS Road, Railway Kodur, Tirupati District, Andhra Pradesh 516101"),
        ("Coordinates / Map", "13.9531717, 79.3515113 (Google Maps integrated)"),
        ("Contact & WhatsApp", "+91 9951021282 | navadeeproy@gmail.com"),
        ("Operating Hours", "8:00 AM – 11:00 PM, All 7 Days (Asia/Kolkata timezone)"),
        ("Fulfillment Model", "STRICTLY Counter Pickup (BOPIS - Buy Online, Pick Up in Store). NO delivery."),
        ("Budget Constraints", "Strictly ₹0 operational budget. Uses free-tier infrastructure only (Firebase Spark, Supabase Free, Groq Free)."),
        ("Pricing Model", "All items sold strictly at Maximum Retail Price (MRP)."),
        ("Language Policy", "English-only site interface (Telugu removed per owner instruction)."),
        ("Printing / Xerox Rates", "B&W ₹2/side, Bulk B&W ₹1.5/side, Double-sided sheet ₹4, ID card copy ₹3, Lamination ₹20, Laminated ID copy ₹40, Colour printing from ₹10/page. No spiral binding.")
    ]
    for k, v in facts:
        p = doc.add_paragraph(style='List Bullet')
        r_k = p.add_run(f"{k}: ")
        r_k.bold = True
        p.add_run(v)

    # Section 2: Complete Technology Stack
    h2 = doc.add_heading("2. Full-Stack Technology Architecture", level=1)
    h2.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    doc.add_paragraph(
        "The project is structured as a modern high-performance single-page application (SPA) backed by serverless cloud infrastructure."
    )

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Tier / Component"
    hdr_cells[1].text = "Technology / Library"
    hdr_cells[2].text = "Role & Implementation Purpose"
    for c in hdr_cells:
        set_cell_background(c, "1F2640")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.bold = True

    stack_rows = [
        ("Frontend Core", "React 18.3 + TypeScript 5.9", "Single Page Application with strict type safety, lazy-loaded route splits."),
        ("Bundler & Dev Tooling", "Vite 8.3", "Ultra-fast ESM bundling with Turbopack-like instant HMR."),
        ("Styling Framework", "Tailwind CSS v4", "Zero-runtime CSS with design token variables and strict color palette."),
        ("Routing", "React Router DOM v6.30", "Client-side routing with scroll reset and path-based Error Boundaries."),
        ("Internationalization", "i18next + react-i18next", "Keyed string translations under src/i18n/en.json."),
        ("3D Engine & Graphics", "Three.js + React Three Fiber + Drei", "Interactive 360° product viewers, single shared WebGL canvas."),
        ("Database (Primary)", "Firebase Firestore v12", "Real-time NoSQL database for categories, products, variants, shop settings."),
        ("Authentication", "Firebase Auth v12", "Anonymous guest login upgraded to Google OAuth upon order submission."),
        ("File Storage", "Supabase Storage v2 (JS client)", "Bucket_Files bucket (Mumbai region) for Xerox documents up to 25MB."),
        ("Serverless & AI", "Supabase Edge Functions + Groq", "Deno runtime calling Groq Llama-3.1-8b-instant for chatbot RAG."),
        ("Admin Seeding", "Firebase Admin SDK v14 (Node 24)", "Direct local seed script with type stripping (npm run seed)."),
        ("Icons", "Lucide React", "Lightweight SVG icon set.")
    ]

    for tier, tech, role in stack_rows:
        row_cells = table.add_row().cells
        row_cells[0].text = tier
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].text = tech
        row_cells[2].text = role
        for c in row_cells:
            set_cell_background(c, "F6F5F2")

    # Section 3: Design System & Tokens
    h3 = doc.add_heading("3. Design System & Token Constraints", level=1)
    h3.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    doc.add_paragraph(
        "The UI follows a calm, paper-like aesthetic reflecting a traditional stationery shop elevated to modern standards. "
        "Strict Rule: NEVER hardcode hex colors in component code. All components must reference semantic design tokens."
    )

    tokens = [
        ("Navy Ink (#1F2640)", "Primary buttons, active navigation, titles, top utility bar."),
        ("Paper White (#FFFFFF)", "Clean background surface for reading and product staging."),
        ("Soft Paper (#F6F5F2)", "Card backgrounds, secondary panels, neutral borders."),
        ("Warm Cream (#F7EFE3)", "Information banners, notice boxes, highlighted sections."),
        ("Powder Blue (#D6E3EF)", "Hero backdrop wall, subtle header accents."),
        ("Desk Sand (#CDBB9E)", "Hero 3D floor plane and subtle grounded dividers."),
        ("Soft Saffron (#EDC27E)", "Product 3D podiums, pedestal radial glow, highlight badges."),
        ("Saffron Accent (#E8952B)", "Accent only: low-stock warnings, 360° badges, primary CTA glow."),
        ("Ink (#1B1F2A)", "Primary body typography text."),
        ("Muted (#6B7280)", "Secondary descriptive text, timestamps, labels."),
        ("Success (#2F8F5B) / Error (#C0453A)", "In-stock indicators, form validations, error notifications.")
    ]
    for name, desc in tokens:
        p = doc.add_paragraph(style='List Bullet')
        r_t = p.add_run(f"{name}: ")
        r_t.bold = True
        p.add_run(desc)

    doc.add_paragraph(
        "Accessibility: Minimum tap targets of 44x44px across all buttons and inputs. Base typography >= 16px. "
        "Two-tone high-contrast focus rings for full keyboard navigation (WCAG 2.2 AA compliant)."
    )

    # Section 4: Current Feature Breakdown
    h4 = doc.add_heading("4. Implemented Features & Application Modules", level=1)
    h4.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    modules = [
        ("Home Page (/) ", [
            "Hero Stage: 3D interactive floating model stage rendered on a single shared WebGL canvas.",
            "Category Grid: 11 category tiles with bespoke pastel tints and category product counters.",
            "Popular Products: Featured curated items with high-resolution transparent cutout photos and hover glow.",
            "How It Works: 3-step visual guide (Browse & Select -> Upload or Cart -> Pick Up at Counter).",
            "Xerox Rate Card: Transparent pricing table with live B&W, colour, and lamination calculations.",
            "Visit Us: Interactive Google Maps iframe, live store status badge (Open Now / Closed calculation using Indian Standard Time), address, phone, directions link."
        ]),
        ("Shop & Catalogue (/shop)", [
            "Complete 48-product catalogue loaded from Firestore (with instant fallback data cache).",
            "Filter Bar: Filter by Category pills, Brand, Paper Ruling, Page Count, and Price Range sliders.",
            "Sorting: Sort by Featured, Price (Low to High), Price (High to Low), and Alphabetical.",
            "Live Search: Client-side search with word stemming, synonym expansion, and IDF relevance scoring.",
            "Pagination: Smooth responsive grid with 'Load More' chunking."
        ]),
        ("Shopping Cart (/cart)", [
            "State Store: Client-side persistent cart in localStorage ('svt-cart-v1').",
            "Quantity Controls: Increment, decrement, remove item, real-time MRP total computation.",
            "Authentication Gate: Forces guest users to complete Google Sign-in before placing orders.",
            "WhatsApp Order Dispatch: Formats complete order breakdown (item names, quantities, variant details, total price) and dispatches directly to the proprietor's WhatsApp."
        ]),
        ("Xerox & Document Portal (/xerox)", [
            "File Uploader: Supports PDF, Word (.doc, .docx), and images (JPEG, PNG, WebP, HEIC). Max 10 files, 25MB per file.",
            "Direct Cloud Upload: Files stream directly to Supabase Storage 'Bucket_Files' without passing through server bottlenecks.",
            "Print Customization: Number of copies, Single vs Double-sided, Black & White vs Colour, ID Card Copy mode, Lamination toggle, customer special instructions.",
            "Order Generation: Generates unique collision-resistant order code (e.g. XR-0922-XXXX) and composes pre-filled WhatsApp message with secure direct download links."
        ]),
        ("AI Shop Assistant Widget", [
            "Floating bottom-right button and slide-up chat drawer.",
            "Tier 1 (Instant Offline Engine): Answers shop FAQs (hours, location, xerox rates, payment options, pickup policy) and searches product catalogue without server latency or cost.",
            "Tier 2 (Serverless RAG): When offline search yields low confidence, queries Supabase Edge Function 'shop-assistant'. Ingests top 60 matching product rows into Groq Llama-3.1-8b-instant with strict anti-hallucination prompts. Disclaimers: 'AI answer. Please confirm prices at counter.' Fallback to WhatsApp if offline."
        ])
    ]

    for title_mod, points in modules:
        h_mod = doc.add_heading(title_mod, level=2)
        h_mod.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)
        for pt in points:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(pt)

    # Section 5: Data Models & Database Architecture
    h5 = doc.add_heading("5. Database Schema & Data Models", level=1)
    h5.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    doc.add_paragraph("Firestore NoSQL structure organized into collections and subcollections:")
    schemas = [
        ("categories/{categoryId}", "{ id: string, name: string, icon: string, tint: string, sortOrder: number }"),
        ("products/{productId}", "{ id: string, name: string, categoryId: string, brand: string, description: string, tags: string[], modelSpec: { model: string, color: string }, image: string, createdAt: timestamp }"),
        ("products/{productId}/variants/{variantId}", "{ id: string, sku: string, name: string, price: number, mrp: number, stock: number, attributes: { pages?: number, ruling?: string, packSize?: number, color?: string } }"),
        ("settings/shop", "{ name: string, proprietor: string, phone: string, email: string, address: string, upiId: string, hours: string, coordinates: { lat: number, lng: number } }"),
        ("orders/{orderId} (Planned)", "{ id: string, orderCode: string, customerUid: string, customerName: string, customerPhone: string, items: Array<{ productId, variantId, name, quantity, price }>, totalAmount: number, paymentStatus: 'pending'|'paid_online'|'pay_at_counter', status: 'received'|'processing'|'ready_for_pickup'|'completed', createdAt: timestamp }")
    ]
    for coll, desc in schemas:
        p = doc.add_paragraph(style='List Bullet')
        r_c = p.add_run(f"{coll}: ")
        r_c.bold = True
        r_c.font.name = 'Consolas'
        p.add_run(desc)

    # Section 6: Security & Guardrails
    h6 = doc.add_heading("6. Security, Rules & Secret Management", level=1)
    h6.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    sec_rules = [
        ("Secret Isolation", "ZERO secrets in git or client bundles. GROQ_API_KEY resides only as a Supabase Edge Function secret. Firebase Service Account is strictly gitignored and stored in web/serviceAccountKey.json."),
        ("Storage Security", "Supabase Bucket_Files configured with INSERT-only policy for anon/authenticated visitors. SELECT/LIST/DELETE policies are disabled so visitors cannot list or wipe customer files."),
        ("Firestore Security Rules", "Public read access to categories, products, and shop settings. Writes restricted to authenticated staff/admin UIDs."),
        ("Edge Function Throttling", "shop-assistant enforces origin whitelisting (localhost, *.web.app, *.firebaseapp.com), 6 requests/min & 40 requests/hour per IP, 200-character input cap, and 15s execution timeout.")
    ]
    for title_s, desc_s in sec_rules:
        p = doc.add_paragraph(style='List Bullet')
        r_s = p.add_run(f"{title_s}: ")
        r_s.bold = True
        p.add_run(desc_s)

    # Section 7: Outstanding Roadmap & Sprint Backlog
    h7 = doc.add_heading("7. Immediate Sprint Backlog & Next Steps", level=1)
    h7.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    backlog = [
        ("Task 1: Full Checkout & Firestore Order Persistence", "Replace pure WhatsApp ordering with direct Firestore order creation (/orders collection), auto-generated sequential/alphanumeric order codes, and client-side UPI Intent deep-links + dynamic QR codes (using shop's UPI ID)."),
        ("Task 2: Product Detail Page (/product/:slug)", "Rich product page featuring interactive 3D model rotation/zoom, photo gallery, variant selector pills, live stock status, and related products carousel."),
        ("Task 3: Live Order Status (/orders/:orderCode)", "Customer order tracking screen with real-time Firestore onSnapshot listener displaying order stage: Received -> Packing / Printing -> Ready for Pickup -> Collected."),
        ("Task 4: Shopkeeper Counter & Admin Apps (/counter & /admin)", "Counter interface for the store operator: real-time incoming order sound alerts, print queue with 1-click PDF download, order status toggle buttons, and inventory stock adjustment."),
        ("Task 5: Supabase Automated File Retention", "Scheduled cleanup script or pg_cron job to delete customer Xerox uploads older than 48 hours to preserve the 1GB free storage limit."),
        ("Task 6: Firebase Hosting Deployment", "Deploy production build to https://sri-vishnu-teja-books-xerox.web.app with security headers, SEO meta tags, and OpenGraph preview cards.")
    ]
    for task_title, task_desc in backlog:
        p = doc.add_paragraph(style='List Bullet')
        r_b = p.add_run(f"{task_title}: ")
        r_b.bold = True
        p.add_run(task_desc)

    # Section 8: Working Rules for ChatGPT
    h8 = doc.add_heading("8. Pair Programming Guidelines for ChatGPT", level=1)
    h8.style.font.color.rgb = RGBColor(0x1F, 0x26, 0x40)

    guidelines = [
        "1. Never suggest paid services, external paid APIs, or paid cloud plans. Always adhere to the ₹0 budget free-tier architecture.",
        "2. Do not introduce delivery, shipping addresses, or shipping fee logic. The store is strictly counter-pickup.",
        "3. Preserve all Tailwind design tokens. Do not introduce arbitrary hex color values in JSX/TSX.",
        "4. Always write TypeScript in strict mode, matching existing interfaces in src/types/models.ts.",
        "5. Keep the single shared WebGL canvas architecture intact. Never mount multiple concurrent full-page 3D canvases.",
        "6. Always keep replies actionable, clean, and modular."
    ]
    for g in guidelines:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(g)

    # Save to memory buffer
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # Repack as .dotm (macroEnabledTemplate)
    zin = zipfile.ZipFile(buf, 'r')
    buf_out = io.BytesIO()
    zout = zipfile.ZipFile(buf_out, 'w', zipfile.ZIP_DEFLATED)

    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == '[Content_Types].xml':
            data = data.replace(
                b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',
                b'application/vnd.ms-word.template.macroEnabledTemplate.main+xml'
            )
        zout.writestr(item, data)
    zout.close()
    zin.close()

    output_path = r"P:\e-commerce sites\PROJECT-CHATGPT-DESCRIPTION.dotm"
    buf_out.seek(0)
    with open(output_path, 'wb') as f:
        f.write(buf_out.read())
    print(f"Generated {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    create_document()
